from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
doc.add_paragraph('Informe técnico de la práctica con Apache Hive', 'Title')
doc.add_paragraph('Universidad del Valle')
doc.add_paragraph('Materia: Tecnologías Emergentes I')
doc.add_paragraph('Estudiante: Ronald Joel Saavedra Vargas')
doc.add_paragraph('Código de práctica: LI16')
doc.add_paragraph('Fecha de inicio: 20 de octubre de 2026')
doc.add_heading('Objetivo', 1)
doc.add_paragraph('Implementar un entorno de Apache Hive y HDFS mediante Docker Compose, cargar el dataset de Facebook y analizar sus datos mediante consultas SQL. Cada apartado incorporará los comandos ejecutados, sus resultados y la evidencia correspondiente.')
doc.add_heading('Entorno de trabajo', 1)
doc.add_paragraph('Se eligió Windows como anfitrión, Ubuntu 24.04 LTS en WSL 2 como entorno de trabajo y Docker Desktop para ejecutar los contenedores Linux. Se acordó crear un proyecto propio con Docker Compose en lugar de clonar Big Data Europe.')
doc.add_heading('Preparación realizada', 1)
doc.add_paragraph('Se verificó la comunicación con Docker Desktop y su motor Linux. La consulta docker ps no mostró contenedores en ejecución. Se instaló Ubuntu 24.04 LTS mediante wsl --install -d Ubuntu-24.04 --no-launch y se confirmó el inicio de Ubuntu 24.04.5 LTS. En Docker Desktop se activó Resources > WSL integration > Ubuntu-24.04 y se aplicó el cambio. Desde Ubuntu se verificaron Docker Engine 29.4.1 y Docker Compose v5.1.3. La consulta docker ps -a identificó tres contenedores detenidos de un proyecto anterior: greenbit-frontend, greenbit-backend y greenbit-db. Se conservaron sin modificaciones.')
doc.add_paragraph('Las salidas de consola verificadas permiten registrar el avance técnico. Las capturas con la hora visible del sistema se incorporarán como evidencia académica antes de cerrar cada punto.')
doc.add_page_break()
doc.add_heading('Desarrollo por puntos', 1)
items = [
('1', 'Revisar los contenedores existentes', 'Verificación inicial realizada: docker ps no mostró contenedores en ejecución. Falta incorporar la captura con la hora del sistema visible.'),
('2', 'Preparar un repositorio diferente de Big Data Europe', 'Decisión tomada: crear un proyecto propio. Pendiente implementar y versionar la configuración.'),
('3', 'Levantar los servicios con Docker Compose', ''),
('4', 'Listar los contenedores del laboratorio', ''),
('5', 'Explicar la función y relación de cada contenedor', ''),
('6', 'Acceder al shell del NameNode', ''),
('7', 'Listar el directorio warehouse y describir lo observado', ''),
('8', 'Validar la dirección y el puerto de la interfaz del NameNode', ''),
('9', 'Explorar warehouse desde la interfaz web de HDFS', ''),
('10', 'Descargar y renombrar el dataset', ''),
('11', 'Preparar una copia del CSV sin encabezado', ''),
('12', 'Copiar el CSV al directorio tmp del NameNode', ''),
('13', 'Subir el archivo a HDFS y verificar su presencia', ''),
('14', 'Acceder al shell del servidor Hive', ''),
('15', 'Crear la base y tabla y cargar los datos', ''),
('16', 'Ejecutar las siete consultas y explicar sus resultados', '')]
for num, title, status in items:
    p = doc.add_paragraph(style='Heading 2')
    p.add_run(f'Punto {num}  {title}').font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph(status or 'Pendiente de ejecución y evidencia.')
doc.save(root / 'informe' / 'Informe_LI16_Ronald_Saavedra.docx')
