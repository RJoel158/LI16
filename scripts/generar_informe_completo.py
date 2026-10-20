from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from contenido_informe import INTRO, POINTS, CONTAINERS

root=Path(__file__).resolve().parents[1]
doc=Document();s=doc.sections[0]
doc.core_properties.author='Ronald Joel Saavedra Vargas'
doc.core_properties.title='Informe técnico LI16 con Apache Hive y HDFS'
s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.08
for style in ['Title','Heading 1','Heading 2']:
 doc.styles[style].font.name='Calibri';doc.styles[style].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(30)
doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 2'].font.size=Pt(13)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('LI16 | Ronald Saavedra | Página ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
s.header.paragraphs[0].add_run('Tecnologías Emergentes I  |  Apache Hive y HDFS').font.size=Pt(9)
doc.add_paragraph('Universidad del Valle')
doc.add_paragraph('Informe técnico de la práctica con Apache Hive','Title')
for line in ['Laboratorio LI16','Estudiante: Ronald Joel Saavedra Vargas','Materia: Tecnologías Emergentes I','Fecha de ejecución: 20 de octubre de 2026']:doc.add_paragraph(line)
doc.add_heading('Resultado del laboratorio',1)
doc.add_paragraph('Se implementó un entorno propio de cuatro contenedores, se cargaron 99.003 usuarios en db_ronald.fb_ronald y se resolvieron las siete consultas requeridas. Los resultados de Hive coinciden con una verificación independiente del CSV original.')
doc.add_paragraph('El desarrollo mantiene la numeración de la guía y destaca cada requerimiento en amarillo. Incluye capturas reales del visor de registros y de la interfaz HDFS, además de imágenes de los registros originales de comandos y SQL. Estas últimas están identificadas como registros renderizados, no como capturas de una terminal; muestran la hora real del sistema al ejecutar el comando.')
doc.add_page_break()
for title,text in INTRO:doc.add_heading(title,1);doc.add_paragraph(text)
doc.add_heading('Trazabilidad de las evidencias',1)
doc.add_paragraph('Los registros completos se conservan en registros como archivos JSON con comando, salida, fecha y código de salida. Las imágenes de registros reproducen sus resultados con formato legible; las advertencias de logging siguen disponibles en el JSON original. No simulan una terminal ni una captura del escritorio. Las capturas de HDFS incluyen un reloj añadido a la interfaz local que lee la hora real del sistema del navegador en UTC-4; este reloj no modifica los datos de HDFS.')

def code(text):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
 run=p.add_run(text);run.font.name='Consolas';run.font.size=Pt(8)
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for rownum,row in enumerate(t.rows):
  for cell in row.cells:
   tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders)
   if rownum==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3');tcPr.append(sh)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
    for run in p.runs:run.font.size=Pt(9);run.bold=(rownum==0)

missing=[]
for key,title,explanation in POINTS:
 p=doc.add_paragraph(style='Heading 1');p.paragraph_format.page_break_before=True;p.add_run('Punto '+title).font.highlight_color=WD_COLOR_INDEX.YELLOW
 if key=='05':
  for name,text in CONTAINERS:doc.add_heading(name,2);doc.add_paragraph(text)
 if key=='15b':
  cols=json.loads((root/'datos/metadata.json').read_text(encoding='utf-8'))['columnas']
  table(['Columna','Tipo Hive'],[(c,'STRING' if c=='gender' else 'BIGINT' if c=='userid' or 'likes' in c else 'INT') for c in cols])
 query=root/'proyecto/sql'/f'{key}.sql'
 if query.exists():doc.add_heading('Consulta ejecutada',2);code(query.read_text(encoding='utf-8'))
 recfile=root/'registros'/f'{key}.json'
 if recfile.exists() and not query.exists():
  rec=json.loads(recfile.read_text(encoding='utf-8-sig'));doc.add_heading('Comando de verificación',2);code(rec['comando'])
 pics=[]
 if key in ('08','09'):pics.append(root/'evidencias'/f'{key}-web.jpg')
 pic=root/'evidencias'/f'{key}.jpg'
 supplied=root/'evidencias'/f'{key}-captura.png'
 if supplied.exists():pics.append(supplied)
 elif pic.exists():pics.append(pic)
 else:
  records=sorted((root/'evidencias').glob(f'{key}-registro-*.png'))
  if key not in ('08','09'):pics.extend(records)
  if not records and key not in ('08','09'):missing.append(key)
 if key=='15e' and (root/'evidencias/15e-web.jpg').exists():pics.append(root/'evidencias/15e-web.jpg')
 doc.add_heading('Evidencia',2)
 for pic in pics:
  if pic.exists():
   with Image.open(pic) as im:width,height=im.size
   maxheight=3.6 if key=='05' else 2.7 if key=='15b' else 5.7 if '-web' in pic.stem and key!='08' else 4.8 if len(pics)>1 else 4.2
   scale=min(6.65/width,maxheight/height)
   doc.add_picture(str(pic),width=Inches(width*scale),height=Inches(height*scale))
   doc.paragraphs[-1].paragraph_format.keep_with_next=True
   kind='Interfaz web real de HDFS' if '-web' in pic.stem else 'Registro real renderizado' if '-registro-' in pic.stem else 'Captura del visor de registros reales'
   p=doc.add_paragraph('Figura '+pic.stem+'  '+kind)
   for run in p.runs:run.font.size=Pt(9)
 if not pics:doc.add_paragraph('Captura pendiente de incorporar. El registro de ejecución está disponible en registros/'+key+'.json.')
 if explanation:doc.add_heading('Interpretación',2);doc.add_paragraph(explanation)

doc.add_page_break();doc.add_heading('Validación de datos y límites del análisis',1)
doc.add_paragraph('Se contrastaron las siete consultas de Hive con cálculos independientes sobre el CSV original. Coinciden los conteos, sumas y promedios. La tabla tiene 99.003 filas, cero identificadores nulos, 175 géneros nulos y dos valores tenure nulos. Las edades van de 13 a 113 años; este rango se describe como dato del archivo, sin verificar que cada edad represente una edad real.')
doc.add_paragraph('Las consultas de género excluyen los valores desconocidos. Los límites de edad siguen los enunciados: jóvenes entre 13 y 25 inclusive, adultos mayores de 25 y el grupo de comparación de likes mayores de 35. Los resultados describen esta muestra histórica, no el comportamiento actual de todos los usuarios de Facebook.')
doc.add_heading('Conclusiones',1)
doc.add_paragraph('El laboratorio mostró la separación entre almacenar datos en HDFS, registrar su estructura en el metastore y consultarlos mediante HiveServer2. La ruta del CSV cambió al cargarlo en la tabla, mientras sus 99.003 registros se conservaron. En este dataset las mujeres tienen más amigos en promedio, los jóvenes reciben más likes en promedio que los mayores de 35 y los dos grupos analizados dan más likes desde celular que desde la web.')
doc.add_heading('Archivos para reproducir la práctica',1)
doc.add_paragraph('El directorio proyecto contiene la configuración Compose y las consultas SQL. datos conserva el ZIP de Kaggle, el CSV original y la copia sin encabezado. registros contiene las salidas completas y evidencias las capturas. scripts contiene la preparación del dataset, la verificación independiente y la generación de este informe. El repositorio de ejecución está en /root/li16-ronald de Ubuntu.')
doc.add_heading('Referencias',1)
for ref in [
'Universidad del Valle. LI16 Práctica con Apache Hive 1. Guía proporcionada para el laboratorio.',
'Apache Hive. Setting Up Hive with Docker. https://hive.apache.org/docs/latest/admin/setting-up-hive-with-docker/',
'Apache Hadoop 3.4.1. HDFS Commands Guide. https://hadoop.apache.org/docs/r3.4.1/hadoop-project-dist/hadoop-hdfs/HDFSCommands.html',
'Sheena Batra. Facebook Data. Kaggle. https://www.kaggle.com/datasets/sheenabatra/facebook-data',
'Guía del dataset de la práctica. https://www.kaggle.com/code/youssefabdelghfar/facebook-data-analysis/input',
]:doc.add_paragraph(ref)
target=root/'informe/Informe_LI16_Ronald_Saavedra.docx';doc.save(target)
print('Documento guardado:',target);print('Puntos sin imagen de evidencia:',', '.join(missing) or 'ninguno')
